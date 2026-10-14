"""
Servidor local para el Reporte de Precios Multi-País (Guatemala, El Salvador, Honduras).

Corre en tu computadora (no en la nube) para ejecutar scripts locales o APIs abiertas
y servir el Dashboard comparativo.

Uso:
    python app.py
    abrir http://localhost:5000 en el navegador
"""

import io
import json
import re
import sys
import time
from pathlib import Path

# Cuando el servidor se arranca con pythonw.exe, redirigir stdout/stderr a servidor.log
if sys.stdout is None or sys.stderr is None:
    _log_file = open(Path(__file__).resolve().parent / "servidor.log", "a", encoding="utf-8", buffering=1)
    sys.stdout = _log_file
    sys.stderr = _log_file

from flask import Flask, jsonify, render_template, request, send_file
import pandas as pd

from matching import build_report, load_catalogo, PAISES_CONFIG, load_config, save_config

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
SCRIPTS_DIR = BASE_DIR / "scripts"
STATUS_FILE = DATA_DIR / "status.json"

DATA_DIR.mkdir(exist_ok=True)

RETAILERS = {
    # Guatemala
    "elektra": {
        "script": SCRIPTS_DIR / "elektra.py",
        "csv": DATA_DIR / "productos_elektra.csv",
        "nombre": "Elektra",
        "pais": "gt",
    },
    "max": {
        "script": SCRIPTS_DIR / "max.py",
        "csv": DATA_DIR / "productos_max.csv",
        "nombre": "MAX",
        "pais": "gt",
    },
    "gallo": {
        "script": SCRIPTS_DIR / "gallo.py",
        "csv": DATA_DIR / "productos_elgallomasgallo.csv",
        "nombre": "El Gallo más Gallo",
        "pais": "gt",
    },
    # El Salvador
    "siman": {
        "script": SCRIPTS_DIR / "siman.py",
        "csv": DATA_DIR / "productos_siman.csv",
        "nombre": "Siman",
        "pais": "sv",
    },
    "curacao_sv": {
        "script": SCRIPTS_DIR / "curacao_sv.py",
        "csv": DATA_DIR / "productos_curacao_sv.csv",
        "nombre": "La Curacao",
        "pais": "sv",
    },
    "prado": {
        "script": SCRIPTS_DIR / "prado.py",
        "csv": DATA_DIR / "productos_prado.csv",
        "nombre": "Prado",
        "pais": "sv",
    },
    # Honduras
    "jetstereo": {
        "script": SCRIPTS_DIR / "jetstereo.py",
        "csv": DATA_DIR / "productos_jetstereo.csv",
        "nombre": "Jetstereo",
        "pais": "hn",
    },
    "elektra_hn": {
        "script": SCRIPTS_DIR / "elektra_hn.py",
        "csv": DATA_DIR / "productos_elektra_hn.csv",
        "nombre": "Elektra HN",
        "pais": "hn",
    },
    "diunsa": {
        "script": SCRIPTS_DIR / "diunsa.py",
        "csv": DATA_DIR / "productos_diunsa.csv",
        "nombre": "Diunsa",
        "pais": "hn",
    },
    # Walmart por país
    "walmart": {
        "script": SCRIPTS_DIR / "walmart_gt.py",
        "csv": DATA_DIR / "productos_walmart_gt.csv",
        "nombre": "Walmart",
        "pais": "gt",
    },
    "walmart_sv": {
        "script": SCRIPTS_DIR / "walmart_sv.py",
        "csv": DATA_DIR / "productos_walmart_sv.csv",
        "nombre": "Walmart",
        "pais": "sv",
    },
    "walmart_hn": {
        "script": SCRIPTS_DIR / "walmart_hn.py",
        "csv": DATA_DIR / "productos_walmart_hn.csv",
        "nombre": "Walmart",
        "pais": "hn",
    },
}

app = Flask(__name__)
app.config["TEMPLATES_AUTO_RELOAD"] = True


def load_status():
    if STATUS_FILE.exists():
        try:
            return json.loads(STATUS_FILE.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            return {}
    return {}


def save_status(status):
    STATUS_FILE.write_text(json.dumps(status, ensure_ascii=False, indent=2), encoding="utf-8")


@app.route("/favicon.ico")
def favicon():
    return "", 204


def get_auth_config():
    p = DATA_DIR / "usuarios.json"
    if not p.exists():
        return None
    try:
        import hashlib
        udata = json.loads(p.read_text(encoding="utf-8"))
        usalt = udata.get("salt", "samsung_retail_2026_ca")
        uhashed = []
        for u in udata.get("usuarios", []):
            ustr = f"{u['usuario'].lower().strip()}:{u['password'].strip()}:{usalt}"
            h = hashlib.sha256(ustr.encode("utf-8")).hexdigest()
            uhashed.append({"usuario": u["usuario"], "nombre": u.get("nombre", u["usuario"]), "hash": h})
        return {"salt": usalt, "users": uhashed}
    except Exception:
        return None

@app.route("/")
def index():
    cfg = get_auth_config()
    return render_template("index.html", auth_config=cfg)


@app.route("/api/config", methods=["GET", "POST"])
def api_config():
    import importlib
    import matching
    importlib.reload(matching)
    
    if request.method == "POST":
        payload = request.get_json(force=True) or {}
        tol = payload.get("tolerancia_pct", payload.get("tolerancia_descuento_pct", 5.0))
        try:
            tol = float(tol)
        except (ValueError, TypeError):
            tol = 5.0
        matching.save_config(DATA_DIR, tol)
        pais = request.args.get("pais", "gt").lower()
        report = matching.build_report(DATA_DIR, tolerancia_pct=tol, pais=pais)
        report["status"] = load_status()
        return jsonify({"ok": True, "tolerancia_pct": tol, "report": report})
    
    tol = matching.load_config(DATA_DIR)
    return jsonify({"tolerancia_pct": tol})


@app.route("/api/data")
def api_data():
    import importlib
    import matching
    importlib.reload(matching)
    tol_param = request.args.get("tolerancia")
    tol = float(tol_param) if tol_param else None
    pais = request.args.get("pais", "gt").lower()
    report = matching.build_report(DATA_DIR, tolerancia_pct=tol, pais=pais)
    report["status"] = load_status()
    return jsonify(report)


def publicar_a_github_sync():
    try:
        import subprocess
        from datetime import datetime
        export_script = BASE_DIR / "exportar_reporte_estatico.py"
        if export_script.exists():
            subprocess.run([sys.executable, str(export_script)], cwd=str(BASE_DIR), capture_output=True, text=True)
        repo_dir = BASE_DIR if (BASE_DIR / ".git").exists() else BASE_DIR.parent / "reporte_web"
        if repo_dir.exists():
            subprocess.run(["git", "add", "-A"], cwd=str(repo_dir), capture_output=True)
            msg = f"Auto-sincronizacion desde dashboard ({datetime.now().strftime('%d/%m/%Y %H:%M')})"
            subprocess.run(["git", "commit", "-m", msg], cwd=str(repo_dir), capture_output=True)
            res = subprocess.run(["git", "push", "origin", "main"], cwd=str(repo_dir), capture_output=True, text=True)
            return res.returncode == 0
    except Exception as e:
        print("[!] Error en publicar_a_github:", e)
    return False

def publicar_a_github_async():
    import threading
    threading.Thread(target=publicar_a_github_sync, daemon=True).start()


@app.route("/api/actualizar/<retailer>", methods=["POST"])
def actualizar(retailer):
    if retailer not in RETAILERS:
        return jsonify({"ok": False, "error": f"Retailer desconocido '{retailer}'"}), 400

    info = RETAILERS[retailer]
    if not info["script"].exists():
        return jsonify({"ok": False, "error": f"No se encontró el script {info['script'].name}"}), 400

    status = load_status()
    actual = status.get(retailer)
    if actual and actual.get("estado") == "corriendo" and time.time() - actual.get("inicio", 0) < 900:
        return jsonify({
            "ok": False,
            "error": f"Ya hay una actualización de {info['nombre']} en curso. Espera a que termine.",
        }), 409

    status[retailer] = {"estado": "corriendo", "inicio": time.time()}
    save_status(status)

    try:
        import subprocess

        result = subprocess.run(
            [sys.executable, str(info["script"])],
            cwd=str(BASE_DIR),
            capture_output=True,
            text=True,
            timeout=900,
        )
        ok = result.returncode == 0 and info["csv"].exists()
        status = load_status()
        if ok:
            status[retailer] = {"estado": "ok", "ultima_actualizacion": time.time()}
            try:
                import threading
                threading.Thread(target=publicar_a_github_async, daemon=True).start()
            except Exception:
                pass
        else:
            log = (result.stdout[-2000:] + "\n" + result.stderr[-2000:]).strip()
            status[retailer] = {"estado": "error", "ultima_actualizacion": time.time(), "log": log}
        save_status(status)

        if not ok:
            return jsonify({"ok": False, "error": "El script terminó con errores", "log": status[retailer].get("log", "")}), 500

        pais = info.get("pais", "gt")
        report = build_report(DATA_DIR, pais=pais)
        try:
            import historico
            historico.registrar_cambios_desde_reporte(report)
        except Exception as e:
            print("[!] Error registrando historico:", e)
        report["status"] = status
        publicar_a_github_async()
        return jsonify({"ok": True, "report": report})

    except Exception as e:
        status = load_status()
        status[retailer] = {"estado": "error", "ultima_actualizacion": time.time(), "log": str(e)}
        save_status(status)
        return jsonify({"ok": False, "error": str(e)}), 500



@app.route("/api/historico", methods=["GET"])
def api_historico():
    try:
        import historico
        metricas = historico.calcular_metricas_historico()
        return jsonify(metricas)
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500

@app.route("/api/historico/exportar-csv", methods=["GET"])
def api_historico_csv():
    try:
        import historico
        csv_path = DATA_DIR / "historico_cambios.csv"
        if not csv_path.exists():
            historico.save_historico(historico.load_historico())
        from flask import send_file
        return send_file(csv_path, mimetype="text/csv", as_attachment=True, download_name="historico_variaciones_precios.csv")
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500


@app.route("/api/publicar-github", methods=["POST"])
def api_publicar_github():
    ok = publicar_a_github_sync()
    if ok:
        return jsonify({"ok": True, "mensaje": "Reporte compilado y publicado exitosamente en GitHub Pages"})
    return jsonify({"ok": False, "error": "No se pudo subir a GitHub. Revisa la conexion o credenciales."}), 500

@app.route("/api/mis-precios", methods=["POST"])
def guardar_mis_precios():
    pais = request.args.get("pais", "gt").lower()
    payload = request.get_json(force=True)
    filas = payload.get("filas", [])
    
    cfg = PAISES_CONFIG.get(pais, PAISES_CONFIG["gt"])
    file_path = DATA_DIR / cfg["mis_precios_file"]
    
    df = pd.DataFrame(filas, columns=["mobile", "memoria", "mi_precio"])
    df.to_csv(file_path, index=False)
    if pais == "gt":
        # Mantener mis_precios.csv sincronizado
        df.to_csv(DATA_DIR / "mis_precios.csv", index=False)

    report = build_report(DATA_DIR, pais=pais)
    report["status"] = load_status()
    return jsonify({"ok": True, "report": report})


def catalogo_editable():
    """Catálogo completo con columnas de los 3 países."""
    path = DATA_DIR / "precios_mercado.csv"
    if not path.exists():
        return []
    df = pd.read_csv(path, dtype=str).fillna("")
    filas = []
    for _, row in df.iterrows():
        mob = str(row.get("MOBILE", "")).strip()
        if not mob:
            continue
        mem = str(row.get("Memoria", "")).strip()
        filas.append({
            "mobile": mob,
            "memoria": mem,
            "rrp_gt": str(row.get("RRP_GT", row.get("RRP", ""))).strip(),
            "msrp_gt": str(row.get("MSRP_GT", row.get("MSRP", ""))).strip(),
            "rrp_sv": str(row.get("RRP_SV", "")).strip(),
            "msrp_sv": str(row.get("MSRP_SV", "")).strip(),
            "rrp_hn": str(row.get("RRP_HN", "")).strip(),
            "msrp_hn": str(row.get("MSRP_HN", "")).strip(),
            # Compatibilidad
            "rrp": str(row.get("RRP_GT", row.get("RRP", ""))).strip(),
            "msrp": str(row.get("MSRP_GT", row.get("MSRP", ""))).strip(),
        })
    return filas


@app.route("/api/catalogo", methods=["GET"])
def api_catalogo_get():
    return jsonify({"filas": catalogo_editable()})


@app.route("/api/catalogo", methods=["POST"])
def api_catalogo_guardar():
    payload = request.get_json(force=True)
    filas = payload.get("filas", [])
    pais = request.args.get("pais", "gt").lower()

    filas_validas = []
    for f in filas:
        mob = str(f.get("mobile", "")).strip()
        if not mob:
            continue
        mem = str(f.get("memoria", "")).strip()
        rrp_gt = str(f.get("rrp_gt", f.get("rrp", ""))).strip()
        msrp_gt = str(f.get("msrp_gt", f.get("msrp", ""))).strip()
        rrp_sv = str(f.get("rrp_sv", "")).strip()
        msrp_sv = str(f.get("msrp_sv", "")).strip()
        rrp_hn = str(f.get("rrp_hn", "")).strip()
        msrp_hn = str(f.get("msrp_hn", "")).strip()

        filas_validas.append({
            "MOBILE": mob,
            "Memoria": mem,
            "RRP_GT": rrp_gt,
            "MSRP_GT": msrp_gt,
            "RRP_SV": rrp_sv,
            "MSRP_SV": msrp_sv,
            "RRP_HN": rrp_hn,
            "MSRP_HN": msrp_hn,
            "RRP": rrp_gt,
            "MSRP": msrp_gt,
        })

    cols = ["MOBILE", "Memoria", "RRP_GT", "MSRP_GT", "RRP_SV", "MSRP_SV", "RRP_HN", "MSRP_HN", "RRP", "MSRP"]
    df = pd.DataFrame(filas_validas, columns=cols)
    df.to_csv(DATA_DIR / "precios_mercado.csv", index=False)

    report = build_report(DATA_DIR, pais=pais)
    report["status"] = load_status()
    publicar_a_github_async()
    return jsonify({"ok": True, "report": report, "catalogo": catalogo_editable()})


@app.route("/api/catalogo/plantilla", methods=["GET"])
def api_catalogo_plantilla():
    """Descarga la plantilla en Excel con columnas para Guatemala, El Salvador y Honduras."""
    path = DATA_DIR / "precios_mercado.csv"
    cols = ["MOBILE", "Memoria", "RRP_GT", "MSRP_GT", "RRP_SV", "MSRP_SV", "RRP_HN", "MSRP_HN"]
    
    if path.exists():
        df = pd.read_csv(path, dtype=str).fillna("")
        for c in cols:
            if c not in df.columns:
                df[c] = ""
        df = df[cols]
    else:
        df = pd.DataFrame(columns=cols)

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Precios Regionales")
        ws = writer.sheets["Precios Regionales"]
        
        from openpyxl.styles import Alignment, Font, PatternFill
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

    output.seek(0)
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="Plantilla_Precios_Samsung_Regional.xlsx"
    )


@app.route("/api/catalogo/importar-excel", methods=["POST"])
def api_catalogo_importar_excel():
    archivo = request.files.get("archivo")
    pais = request.args.get("pais", "gt").lower()
    if not archivo or not archivo.filename:
        return jsonify({"ok": False, "error": "No se recibió ningún archivo."}), 400

    try:
        df = pd.read_excel(archivo)
    except Exception as e:
        return jsonify({"ok": False, "error": f"No se pudo leer el archivo Excel: {e}"}), 400

    # Diccionario de columnas normalizadas
    def match_column(*aliases):
        for col in df.columns:
            cleaned = re.sub(r"[^a-zA-Z0-9]", "", str(col).lower()).strip()
            for alias in aliases:
                target = re.sub(r"[^a-zA-Z0-9]", "", alias.lower()).strip()
                if cleaned == target or cleaned.startswith(target):
                    return col
        return None

    col_mob = match_column("mobile", "modelo", "producto", "model", "celular")
    col_mem = match_column("memoria", "memory", "ram", "capacidad", "rom")
    
    if not col_mob or not col_mem:
        return jsonify({
            "ok": False,
            "error": (
                "El Excel debe contener al menos las columnas 'MOBILE' (o 'Modelo') y 'Memoria'. "
                f"Columnas detectadas: {', '.join(str(c) for c in df.columns)}"
            ),
        }), 400

    # Mapeo de columnas por país
    col_rrp_gt = match_column("rrpgt", "rrpguatemala", "rrpq", "rrp")
    col_msrp_gt = match_column("msrpgt", "msrpguatemala", "msrpq", "msrp")
    
    col_rrp_sv = match_column("rrpsv", "rrpsalvador", "rrpusd", "rrpdolar")
    col_msrp_sv = match_column("msrpsv", "msrpsalvador", "msrpusd", "msrpdolar")
    
    col_rrp_hn = match_column("rrphn", "rrphonduras", "rrphnl", "rrplempira")
    col_msrp_hn = match_column("msrphn", "msrphonduras", "msrphnl", "msrplempira")

    # Si ya existe catálogo previo, lo leemos para conservar valores existentes de columnas no provistas
    pre_existente = {}
    path_mercado = DATA_DIR / "precios_mercado.csv"
    if path_mercado.exists():
        try:
            df_prev = pd.read_csv(path_mercado, dtype=str).fillna("")
            for _, r in df_prev.iterrows():
                key = (str(r.get("MOBILE", "")).strip().upper(), str(r.get("Memoria", "")).strip().upper())
                pre_existente[key] = r
        except Exception:
            pass

    filas_nuevas = []
    for _, row in df.iterrows():
        mob = str(row[col_mob]).strip() if pd.notna(row[col_mob]) else ""
        if not mob or mob.lower() == "nan":
            continue
        mem = str(row[col_mem]).strip() if pd.notna(row[col_mem]) else ""
        key = (mob.upper(), mem.upper())
        prev = pre_existente.get(key, {})

        def val_or_prev(col_match, prev_key):
            if col_match and pd.notna(row.get(col_match)):
                val = str(row[col_match]).strip()
                if val.lower() != "nan" and val != "":
                    return val
            if isinstance(prev, dict) or isinstance(prev, pd.Series):
                return str(prev.get(prev_key, "")).strip()
            return ""

        rrp_gt_val = val_or_prev(col_rrp_gt, "RRP_GT") or val_or_prev(col_rrp_gt, "RRP")
        msrp_gt_val = val_or_prev(col_msrp_gt, "MSRP_GT") or val_or_prev(col_msrp_gt, "MSRP")
        rrp_sv_val = val_or_prev(col_rrp_sv, "RRP_SV")
        msrp_sv_val = val_or_prev(col_msrp_sv, "MSRP_SV")
        rrp_hn_val = val_or_prev(col_rrp_hn, "RRP_HN")
        msrp_hn_val = val_or_prev(col_msrp_hn, "MSRP_HN")

        filas_nuevas.append({
            "MOBILE": mob,
            "Memoria": mem,
            "RRP_GT": rrp_gt_val,
            "MSRP_GT": msrp_gt_val,
            "RRP_SV": rrp_sv_val,
            "MSRP_SV": msrp_sv_val,
            "RRP_HN": rrp_hn_val,
            "MSRP_HN": msrp_hn_val,
            "RRP": rrp_gt_val,
            "MSRP": msrp_gt_val,
        })

    cols = ["MOBILE", "Memoria", "RRP_GT", "MSRP_GT", "RRP_SV", "MSRP_SV", "RRP_HN", "MSRP_HN", "RRP", "MSRP"]
    nuevo_df = pd.DataFrame(filas_nuevas, columns=cols)
    nuevo_df.to_csv(path_mercado, index=False)

    report = build_report(DATA_DIR, pais=pais)
    report["status"] = load_status()
    return jsonify({
        "ok": True,
        "report": report,
        "catalogo": catalogo_editable(),
        "filas_importadas": len(nuevo_df),
        "columnas_detectadas": {
            "GT": f"RRP={col_rrp_gt or 'None'}, MSRP={col_msrp_gt or 'None'}",
            "SV": f"RRP={col_rrp_sv or 'None'}, MSRP={col_msrp_sv or 'None'}",
            "HN": f"RRP={col_rrp_hn or 'None'}, MSRP={col_msrp_hn or 'None'}",
        }
    })


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, threaded=True, use_reloader=False)